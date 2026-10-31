from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from datetime import datetime,timedelta,timezone
from pathlib import Path
import json,subprocess
root=Path(__file__).resolve().parents[1]/'web/assets'
headers=['날짜','수업','구간','제목','운동 내용','기록 종류','관련 동작','오늘 운동 설명 영상','수업 시간(공휴일)']
movements=['Front Squat','Back Squat','Deadlift','Clean','Clean & Jerk','Snatch','Power Snatch','Squat Snatch','Strict Press','Push Press','Split Jerk','Row','Box Step Over','Pull-up','Ring Dip','Burpee']
catalogue=json.loads(subprocess.check_output(['node','--input-type=module','-e',"import {DEFAULT_RM_CATALOGUE} from './web/rm-catalogue.js';console.log(JSON.stringify(DEFAULT_RM_CATALOGUE.map(m=>m.name)))"],cwd=root.parents[1],text=True))
movements=list(dict.fromkeys(catalogue+movements))
now=datetime.now(timezone(timedelta(hours=9))).date(); monday=now+timedelta(days=(7-now.weekday())%7 or 7)
examples=[[str(monday),'B/T','S1','Front Squat 1RM','워밍업 후 1회 최대 중량 측정','1RM','Front Squat',''],[str(monday),'B/T','S2','5 Sets · Every 3:30','8 Front Squat (115 / 80 lbs)\n16 Box Step Over\n8 Front Squat','중량','Front Squat; Box Step Over',''],[str(monday),'WOD','S1','Snatch Complex · 5 Sets','1 Snatch Lift Off\n1 Snatch Deadlift\n1 Snatch Pull\n1 Squat Snatch','중량','Snatch',''],[str(monday),'WOD','S2','6 Sets · Every 1:30 / Rest 1:00','ODD: 20/18 cal Row + Max Power Snatch\nEVEN: 10 Power Snatch + Max cal Row','세트별 횟수·칼로리','Row; Power Snatch','']]
for filled,name in [(False,'onoff-workouts-template.xlsx'),(True,'onoff-workouts-example.xlsx')]:
 w=Workbook();s=w.active;s.title='수업등록';s.append(headers)
 if filled:
  for row in examples:s.append(row)
 for i,width in enumerate([15,12,10,32,65,24,34,42,28],1):s.column_dimensions[chr(64+i)].width=width
 s.freeze_panes='A2';s.auto_filter.ref='A1:I1001'
 for c in s[1]:c.fill=PatternFill('solid',fgColor='163564');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(vertical='center')
 s.row_dimensions[1].height=28
 for row in s.iter_rows(min_row=2,max_row=max(10,s.max_row),max_col=9):
  for c in row:c.alignment=Alignment(wrap_text=True,vertical='top');c.number_format='@'
  s.row_dimensions[row[0].row].height=75 if filled else 24
 for col,formula in [('B','"B/T,WOD"'),('F','"없음,중량,1RM,시간,횟수,라운드 + 횟수,칼로리,세트별 횟수·칼로리,완료 여부"')]:
  dv=DataValidation(type='list',formula1=formula,allow_blank=False);dv.errorTitle='항목 확인';dv.error='목록에서 선택해주세요.';dv.showErrorMessage=True;s.add_data_validation(dv);dv.add(f'{col}2:{col}1001')
 ex=w.create_sheet('작성예시');ex.append(headers)
 for row in examples:ex.append(row)
 for i,width in enumerate([15,12,10,32,65,24,34,42,28],1):ex.column_dimensions[chr(64+i)].width=width
 for row in ex:
  for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
  ex.row_dimensions[row[0].row].height=75
 instructions=w.create_sheet('작성안내');instructions.column_dimensions['A'].width=110
 for line in ['ONOFF CROSSFIT · 수업 일괄 등록','1. 수업등록 시트에만 입력하세요. 작성예시 시트는 업로드되지 않습니다.','2. 날짜: YYYY-MM-DD. 한 행은 한 수업의 한 구간입니다. S1부터 순서대로 작성하세요.','3. 수업: B/T 또는 WOD. 수요일은 WOD만, 토·일요일은 기본 휴강입니다. 공휴일 특별 수업은 수업 시간을 지정하세요.','4. 운동 내용은 셀 안에서 줄바꿈(Alt+Enter)할 수 있습니다. 기본 중량 단위는 lbs입니다.','5. 기록 종류는 목록에서 선택하세요. 기록이 필요 없으면 없음을 선택하세요.','6. 관련 동작은 운동 사전의 정확한 이름으로 입력하고, 여러 개는 세미콜론(;)으로 구분하세요.','7. 1RM 구간은 관련 동작을 정확히 하나 입력해야 개인 1RM과 연결됩니다.','8. 동작 설명 영상은 앱의 운동 사전에 한 번 등록합니다. 엑셀에 반복 입력하지 않아도 됩니다.','9. 오늘 운동 설명 영상은 선택 항목입니다. 해당 수업 전용 유튜브 링크를 넣으세요.','10. 수업은 해당 날짜 전날 운영자가 설정한 한국 시각에 자동 공개됩니다. 기본값은 22:00입니다. 운영자는 미리 볼 수 있습니다.','11. 업로드 미리보기에서 날짜와 오류를 확인한 뒤 저장하세요. 기존 수업 덮어쓰기는 별도 확인합니다.','공휴일 특별 수업: 수업 시간(공휴일) 열에 10:00, 11:00처럼 시작 시각을 입력합니다. 각 수업은 1시간입니다.','추가 종목: 운영자가 앱에서 등록한 최신 종목 이름을 사용할 수 있습니다. 동작목록 시트는 기본 목록 참고용입니다.', '12. 예시파일은 기능 확인용입니다. 날짜와 내용을 실제 수업으로 바꾼 뒤 등록하세요.']:
  instructions.append([line]);instructions.cell(instructions.max_row,1).alignment=Alignment(wrap_text=True,vertical='top');instructions.row_dimensions[instructions.max_row].height=32
 ms=w.create_sheet('동작목록');ms.append(['관련 동작 이름']);[ms.append([m]) for m in movements];ms.column_dimensions['A'].width=35
 w.save(root/name)
print('Created Excel template and separate example workbook.')
